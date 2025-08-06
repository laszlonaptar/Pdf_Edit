from fastapi import FastAPI, Request, Form
from fastapi.responses import FileResponse, HTMLResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from datetime import datetime
from openpyxl import load_workbook
import uuid
import os

app = FastAPI()

app.mount("/static", StaticFiles(directory="static"), name="static")
templates = Jinja2Templates(directory="templates")

@app.get("/", response_class=HTMLResponse)
async def serve_form(request: Request):
    return templates.TemplateResponse("index.html", {"request": request})

@app.post("/generate_excel")
async def generate_excel(
    request: Request,
    datum: str = Form(...),
    bauort: str = Form(...),
    bf: str = Form(...),
    taetigkeit: str = Form(...),
    beginn: str = Form(...),
    ende: str = Form(...),
    geraet: str = Form(""),
    name1: str = Form(...),
    vorname1: str = Form(...),
    ausweis1: str = Form(...),
):
    # Munkaidő számítás
    start = datetime.strptime(beginn, "%H:%M")
    end = datetime.strptime(ende, "%H:%M")
    duration = (end - start).seconds / 3600

    # Automatikus szünetlevonás
    pause = 0
    if start <= datetime.strptime("09:00", "%H:%M") < end:
        pause += 0.25
    if start <= datetime.strptime("12:00", "%H:%M") < end:
        pause += 0.75

    work_hours = duration - pause

    # Excel sablon betöltés
    wb = load_workbook("GP-t.xlsx")
    ws = wb.active

    ws["D4"] = datum
    ws["D5"] = bauort
    ws["D6"] = bf
    ws["B8"] = taetigkeit
    ws["B12"] = name1
    ws["C12"] = vorname1
    ws["D12"] = ausweis1
    ws["F12"] = beginn
    ws["G12"] = ende
    ws["H12"] = work_hours
    ws["I12"] = work_hours
    ws["B10"] = geraet

    # Fájl mentése
    filename = f"arbeitsnachweis_{uuid.uuid4().hex[:8]}.xlsx"
    filepath = f"/tmp/{filename}"
    wb.save(filepath)

    return FileResponse(filepath, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename=filename)
